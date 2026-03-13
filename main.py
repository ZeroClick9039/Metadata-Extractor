import os
import customtkinter as ctk
from tkinterdnd2 import DND_FILES, TkinterDnD
from tkinter import messagebox
import magic
import piexif                   # FOR IMAGES JPEG AND JPG
from PyPDF2 import PdfReader    # FOR PDF
import pikepdf
from docx import Document       # FOR docx
from openpyxl import load_workbook  # For XLSX
from pptx import Presentation       # For PPTX 


# Set appearance and theme
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

class Application(ctk.CTkFrame):
        def __init__(self, master):
                super().__init__(master)
                self.grid(row=0, column=0, sticky="nsew")
                
                master.title("Metadata Extractor")
                master.geometry("600x400")
                master.minsize(700, 400)
                self.file_path = None
                
# ===== Configure grid ======
                master.grid_rowconfigure(0, weight=1)
                master.grid_columnconfigure(0, weight=1)

# ======= Main Layout Grid ======
                self.grid_rowconfigure(0, weight=1)
                self.grid_columnconfigure(0, weight=1)
                # self.grid_columnconfigure(1, weight=1)
                
# ======== Tab for extracting and editing metadata =============               
                
                self.my_tab = ctk.CTkTabview(self)
                self.my_tab.grid(row=0, column=0, sticky="nsew")
                
                self.show_tab = self.my_tab.add("Show")
                self.edit_tab = self.my_tab.add("Edit")
                
                self.show_tab.grid_rowconfigure(0, weight=1)
                self.show_tab.grid_columnconfigure(1, weight=1)
                self.show_tab.grid_rowconfigure(0, weight=1)
                self.edit_tab.grid_rowconfigure(0, weight=1)
                self.edit_tab.grid_columnconfigure(1, weight=1)
                self.edit_tab.grid_rowconfigure(0,weight=1)
                
                
                               
# ======= Left Panel =======
                self.left_panel = ctk.CTkFrame(self.show_tab, width=200,height=200, corner_radius=0, fg_color="#010104")
                self.left_panel.grid(row=0, column=0, sticky="ns")
                self.left_panel.grid_rowconfigure(0, weight=1)
                self.left_panel.grid_columnconfigure(0, weight=1)
                
        # ===== Title Label in Sidebar =======
                self.sidebar_title = ctk.CTkLabel(
                        self.left_panel,
                        text ="📂 File Drop Area",
                        font=ctk.CTkFont(size=18, weight="bold"),
                        anchor="center"
                )
                self.sidebar_title.grid(row=0, column=0, padx=10, pady=20, sticky="news")
                
                
        # ======= Enable drag & drop ========
                self.sidebar_title.drop_target_register(DND_FILES)
                self.sidebar_title.dnd_bind("<<Drop>>", lambda e: self.drop_file(e, self.file_display))
                
        # ======= Display file path =========
                self.file_display = ctk.CTkTextbox(self.left_panel, height=120, width=200)
                self.file_display.grid(row=2, column=0, padx=20, pady=20, sticky="ew")
                self.file_display.insert("1.0", "No file dropped yet...")
                self.file_display.configure(state="disabled")
                
        # ======= Extact Metadata button =========        
                self.extract_btn = ctk.CTkButton(self.left_panel, text="Extract Metadata",command=self.metadata)
                self.extract_btn.grid(row=3, column=0, padx=10,pady=20, sticky="ew")
                
                self.edit_btn = ctk.CTkButton(self.left_panel, text="Save ",command= lambda: self.save_custom_metadata(self.file_display))
                self.edit_btn.grid(row=4, column=0, pady=10, padx=2, sticky="ew")
                
# ======== Right Panel ======== 
                self.right_panel = ctk.CTkFrame(self.show_tab, corner_radius=10, fg_color="#080811")
                self.right_panel.grid(row=0, column=1, sticky="nsew", padx=15, pady=15)
                self.right_panel.grid_rowconfigure(0, weight=0)
                self.right_panel.grid_rowconfigure(1, weight=1)
                self.right_panel.grid_columnconfigure(0, weight=1)
                
                self.topbar = ctk.CTkFrame(self.right_panel, fg_color="transparent")
                self.topbar.grid(row=0, column=0, sticky="ew", padx=6, pady=6)
                self.topbar.grid_columnconfigure(0, weight=1)
                self.topbar.grid_columnconfigure(1, weight=0)
                
                self.copy_status = ctk.CTkLabel(self.topbar, text="", fg_color="transparent")
                self.copy_status.grid(row=0, column=0, sticky="w", padx=(6, 0))
                
        # ======= copy button to copy metadata =========
                self.copy_btn = ctk.CTkButton(self.topbar, text="Copy",width=80, command=self.copy_metadata)
                self.copy_btn.grid(row=0,column=1,sticky='e',padx=(0,6))
                
        # ======= Text area (Metadata area) =========
                self.display_metadata = ctk.CTkTextbox(self.right_panel, wrap="word")
                self.display_metadata.grid(row=1,column=0,padx=10,pady=6,sticky="nsew")

# ============== Edit Tab ================
                self.edit_left_panel = ctk.CTkFrame(self.edit_tab, width=60, height=600, corner_radius=0, fg_color="#010104")
                self.edit_left_panel.grid(row=0, column=0, sticky="ns")
                self.edit_left_panel.grid_rowconfigure(0, weight=1)
                self.edit_left_panel.grid_columnconfigure(0, weight=1)
        # ============= Title Label in Sidebar ==================
                self.edit_sidebar_title = ctk.CTkLabel(
                        self.edit_left_panel,
                        text="📂 File Drop Area",
                        font=ctk.CTkFont(size=18, weight="bold"),
                        anchor="center"
                )
                self.edit_sidebar_title.grid(row=0, column=0, padx=10, pady=20, sticky="nsew")
                
        # ========== Enable drag & drop =============
                self.edit_sidebar_title.drop_target_register(DND_FILES)
                self.edit_sidebar_title.dnd_bind("<<Drop>>",lambda e: self.drop_file(e, self.edit_file_display))
        
        # ========== Display file path =============
                self.edit_file_display = ctk.CTkTextbox(self.edit_left_panel, height=120, width=200)
                self.edit_file_display.grid(row=2, column=0, padx=20, pady=20, sticky="ew")
                self.edit_file_display.insert("1.0", "No file dropped yet...")
                self.edit_file_display.configure(state="disabled")
        
        # ========== Extract Metadata Button ==========
                self.edit_btn = ctk.CTkButton(self.edit_left_panel, text="show Metadata")
                self.edit_btn.grid(row=3, column=0, padx=10, pady=20, sticky="ew")
                
# ============ edit Right Panel ==============
                self.edit_right_panel = ctk.CTkFrame(self.edit_tab, corner_radius=10, fg_color="#080811")
                self.edit_right_panel.grid(row=0, column=1, sticky="nsew", padx=15, pady=15)
                self.edit_right_panel.grid_rowconfigure(0, weight=0)
                self.edit_right_panel.grid_rowconfigure(0, weight=1)
                self.edit_right_panel.grid_columnconfigure(0, weight=1)
                
                # self.topbar = ctk.CTkFrame(self.edit_right_panel, fg_color="transparent")
                # self.topbar.grid(row=0, column=0, sticky="ew", padx=6, pady=6)
                # self.topbar.grid_columnconfigure(0, weight=1)
                # self.topbar.grid_columnconfigure(1, weight=0)
        
                self.edit_display_metadata = ctk.CTkTextbox(self.edit_right_panel, wrap="word")
                self.edit_display_metadata.grid(row=0, column=0, padx=10, pady=6, sticky="nsew")
                
                
        def metadata(self):
                if not self.file_path:
                        messagebox.showwarning("File error","No file selected")
                        return
                
                filename = os.path.basename(self.file_path)
                _, ext = os.path.splitext(filename)
                ext = ext.lower()
                try:
                        mime = magic.from_file(self.file_path, mime=True)
                except Exception as e:
                        print(f"[!] Magic failed: {e}")
                        mime = None
                
                if mime and mime.startswith("image/"):
                        self.image_metadata()
                elif (mime in ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","application/vnd.ms-excel"] or ext in [".xlsx",".xls"]) :
                        self.xlsx_metadata()
                elif mime == "application/pdf" or ext == ".pdf":
                        self.pdf_metadtata()
                elif (mime in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document","application/msword"] or ext in [".docx", ".doc"] or mime == "application/zip"):
                        self.docx_metadata()
                elif (mime in ["application/vnd.openxmlformats-officedocument.presentationml.presentation","application/vnd.ms-powerpoint"] or ext in[".pptx",".ppt"]):
                        self.ppt_metadata()
                else:
                        messagebox.showwarning("Unsupported file", "Enter a compatible file")
                
        
        def drop_file(self, event, display_widget=None):
                
                self.file_path = event.data.strip("{}")
                
                if display_widget is None:
                        display_widget = self.file_display
                        
                display_widget.configure(state="normal")
                display_widget.delete("1.0", "end")
                
                if os.path.exists(self.file_path):
                        display_widget.insert("1.0", f"File Dropped:\n{self.file_path}")
                else:
                        display_widget.insert("1.0", "Invalid File")
                
                display_widget.configure(state="disabled")
                        
        def friendly_value(self,value):
                if isinstance(value, bytes):
                        return value.decode(errors="ignore")
                elif isinstance(value, tuple):
                        if len(value) == 2 and all(isinstance(x,int) for x in value):
                                num, den = value
                                return num/den if den != 0 else None
                        return tuple(self.friendly_value(v) for v in value)
                return value
        
        def coordinates(self,dms_tuple):
                try: 
                        degree = self.friendly_value(dms_tuple[0]) or 0
                        minute = self.friendly_value(dms_tuple[1]) or 0
                        second = self.friendly_value(dms_tuple[2]) or 0
                        return degree+(minute/60)+(second/3600)
                except Exception:
                        return None
        
        def gps_time_to_str(self, timestamp_tuple):
                try:
                        h, m, s = (int(self.friendly_value(x) or 0) for x in timestamp_tuple)
                        return f"{h:02d}:{m:02d}:{s:02d} UTC"
                except Exception:
                        return None
                        
        def image_metadata(self):
                if not self.file_path:
                        messagebox.showwarning("File Error","No file selected")
                
                try:
                        exif = piexif.load(self.file_path)
                except Exception as e:
                        self.display_metadata.delete("1.0", "end")
                        self.display_metadata.insert("end", f"Error loading metadata:\n{e}")
                        return
                
                summary_lines = []
                zeroth = exif.get("0th",{})
                for tag_id, val in zeroth.items():
                        name = piexif.TAGS["0th"].get(tag_id, {}).get("name", str(tag_id))
                        if name in ("ImageWidth", "ImageLength", "Make", "Model", "DateTime"):
                                summary_lines.append(f"{name}: {self.friendly_value(val)}")
                                
                gps_ifd = exif.get("GPS",{})
                if gps_ifd:
                        lat = None; lon = None
                        for tag_id, val in gps_ifd.items():
                                tag_name = piexif.TAGS["GPS"].get(tag_id, {}).get("name", str(tag_id))
                                if tag_name == "GPSLatitude":
                                        lat = self.coordinates(val)
                                elif tag_name == "GPSLongitude":
                                        lon = self.coordinates(val)
                        if lat is not None and lon is not None:
                                summary_lines.append(f"GPS: {lat:.6f},{lon:6f}")
                
                self.display_metadata.delete("1.0", "end")
                                        
                if summary_lines:
                        self.display_metadata.insert("end", "---- Summary ----\n")
                        for L in summary_lines:
                                self.display_metadata.insert("end", f"{L}\n")
                        self.display_metadata.insert("end", "\n")
                        
                for ifd_name in ("0th", "Exif", "GPS","Interop","1st"):
                        ifd = exif.get(ifd_name)
                        if not ifd:
                                continue
                        self.display_metadata.insert("end",f"\n--- {ifd_name} ---\n")
                        for tag_id, val in ifd.items():
                                tag_entry = piexif.TAGS[ifd_name].get(tag_id, {})
                                tag_name = tag_entry.get("name", str(tag_id))
                                
                                if tag_name == "PixelXDimension":
                                        tag_name = "ExifImageWidth"
                                elif tag_name == "PixelYDimension":
                                        tag_name = "ExifImageHeight"
                                elif tag_name == "GPSLatitude":
                                        val = self.coordinates(val)
                                elif tag_name == "GPSLongitude":
                                        val = self.coordinates(val)
                                elif tag_name == "GPSTimeStamp":
                                        val = self.gps_time_to_str(val)
                                        
                                self.display_metadata.insert("end",f"{tag_name}:       {self.friendly_value(val)}\n")
                        self.display_metadata.insert("end","\n")
                
                                        
                self.copy_status.configure(text="")
                            
        def copy_metadata(self):
                data = self.display_metadata.get("1.0", "end-1c")
                
                if not data.strip():
                        self.copy_status.configure(text="Nothing to Copy")
                        self.after(1600, lambda: self.copy_status.configure(text=""))
                        return
                
                try:
                        self.clipboard_clear()
                        self.clipboard_append(data)
                        self.update()
                        self.copy_status.configure(text="Copied ✓")
                        
                        self.after(1200, lambda: self.copy_status.configure(text=""))
                except Exception as e:
                        messagebox.showerror("Copy failed", f"Could not copy metadata:\n{e}")
                        
        def pdf_metadtata(self):
                if not self.file_path:
                        messagebox.showwarning("File Error","No file selected")
                
                try:
                        file = PdfReader(self.file_path)
                        info = file.metadata
                        
                        self.display_metadata.delete("1.0", "end")
                        self.display_metadata.insert("end","----- PDf Metadata =====\n")
                        if info:
                                for key, value in info.items():
                                        
                                        try:
                                                val = value.get_object() if hasattr(value, "get_object") else value
                                        except Exception:
                                                val = value
                                                
                                        self.display_metadata.insert("end",f"{key}: {val}\n")
                        else:
                                self.display_metadata.insert("end", "No metadata found.\n")
                        
                        self.copy_status.configure(text="")
                except Exception as e:
                        messagebox.showerror("Error", f"Failed to read Pdf metadata:\n{e}")
                        
        # def get_metadata_from_textbox(self, textbox_widget):
        #         metadata = {}
        #         text = textbox_widget.get("1.0", "end").strip()
        #         for line in text.splitlines():
        #                 if ":" in line:
        #                         key, value = line.split(":", 1)
        #                         metadata[key.strip()] = value.strip()
        #         return metadata

        def save_custom_metadata(self, textbox_widget):
                # Collect metadata from textbox
                new_metadata = {}
                content = textbox_widget.get("1.0", "end-1c").strip()  # avoid trailing newline
                
                for line in content.splitlines():
                        if ":" in line:
                                key, value = line.split(":", 1)
                                key = key.strip()
                                value = value.strip()
                                if not key.startswith("/"):
                                        key = "/" + key  # ensure proper PDF metadata key format
                                new_metadata[key] = value

                # Open and update PDF safely
                with pikepdf.open(self.file_path) as pdf:
                        pdf.docinfo.clear()  # delete all old metadata
                        for key, value in new_metadata.items():
                                pdf.docinfo[key] = value

                        # Save with new name
                        output_path = self.file_path.replace(".pdf", "_updated.pdf")
                        pdf.save(output_path)

                return output_path  # return the new file path


        
        # def update_pdf_metadata(self):
        #         if not self.file_path or not self.file_path.lower().endswith(".pdf"):
        #                 messagebox.showwarning("File Errors", "Please drop a PDF file first")
        #                 return
                
        #         new_metadata = self.get_metadata_from_textbox(self.display_metadata)
                
        #         try:
        #                 pdf = pikepdf.open(self.file_path)
                        
        #                 for key, value in new_metadata.items():
        #                         if not key.startswith("/"):
        #                                 key = "/" + key
        #                         pdf.docinfo[key] = value
                                
        #                 output_path = self.file_path.replace(".pdf", "_updataed.pdf")
        #                 pdf.save(output_path)
        #                 pdf.close()
                        
        #                 messagebox.showinfo("Success", f"PDF metadata updated!\n Saved as: {output_path}")
                
        #         except Exception as e:
        #                 messagebox.showerror("Error", f"failed to update PDF metadata:\n{e}")
        
                        
        def docx_metadata(self):
                if not self.file_path:
                        messagebox.showwarning("File Error", "No file selected")
                
                try:
                        file = Document(self.file_path)
                        
                        props =  file.core_properties
                        
                        names = ["Tile","Subject","author","keywords","comment","last_modified_by","revision","category","content_status","identifier","language","version","created","modified","last_printed"]
                        
                        self.display_metadata.delete("1.0", "end")
                        
                        for name in names:
                                value = getattr(props, name, None)
                                if value is not None:
                                        self.display_metadata.insert("end",f"{name}: {value}\n")
                                else:
                                        self.display_metadata.insert("end", f"{name}: None\n")
                                        
                except Exception as e:
                        messagebox.showerror("Error", f"Failed to read Pdf metadata:\n{e}")
                        
        def xlsx_metadata(self):
                if not self.file_path:
                        messagebox.showwarning("File Error", "No file selected")
                try:
                        file = load_workbook(self.file_path)
                        
                        props = file.properties
                        names = ["title","creator","description","subject","identifier","language","created","modified","lastModifiedBy","category","contentStatus","revision","keywords","lastPrinted"]
                        
                        self.display_metadata.delete("1.0", "end")
                        
                        for name in names:
                                value = getattr(props, name, None)
                                if value is not None:
                                        self.display_metadata.insert("end",f"{name}: {value}\n")
                                else:
                                        self.display_metadata.insert("end", f"{name}: None\n")
                        
                except Exception as e:
                        messagebox.showerror("Error", f"Failed to read xlsx metadata:\n{e}")
        
        def ppt_metadata(self):
                if not self.file_path:
                        messagebox.showwarning("File Error", "No file selected")
                try:
                        prs = Presentation(self.file_path)
                        props = prs.core_properties
                        
                        names = ["title","author","subjects","keywords","created","last_modified_by","props.revision"]
                        
                        self.display_metadata.delete("1.0", "end")
                        
                        for name in names:
                                value = getattr(props, name,None)
                                if value is not None:
                                        self.display_metadata.insert("end",f"{name}: {value}\n")
                                else:
                                        self.display_metadata.insert("end", f"{name}: None\n")
                except Exception as e:
                        messagebox.showerror("Error", f"Failed to read pptx metadata:\n{e}")
                        
        

if __name__ == "__main__":
        root = TkinterDnD.Tk()
        app = Application(root)
        app.mainloop()
                